"""
CloudNexus RAG Service

Pipeline:

PDF / DOCX / PPTX / TXT / MD
        ↓
Text extraction
        ↓
Chunking
        ↓
Gemini embeddings
        ↓
ChromaDB
        +
BM25 keyword retrieval
        ↓
Reciprocal Rank Fusion (RRF)
        ↓
Top-K relevant chunks
"""

import os
import re
import chromadb 

from rapidfuzz.fuzz import partial_ratio 
from pathlib import Path
from typing import Any

from dotenv import load_dotenv
from rank_bm25 import BM25Okapi

from langchain_core.documents import Document
from langchain_community.document_loaders import (PyPDFLoader,Docx2txtLoader,TextLoader,)
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_google_genai import GoogleGenerativeAIEmbeddings


# ============================================================
# ENVIRONMENT
# ============================================================

load_dotenv()

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")

if not GEMINI_API_KEY:
    raise RuntimeError(
        "GEMINI_API_KEY is missing from .env"
    )


# ============================================================
# EMBEDDINGS
# ============================================================

embeddings = GoogleGenerativeAIEmbeddings(
    model="gemini-embedding-001",
    google_api_key=GEMINI_API_KEY,
)


# ============================================================
# CHROMA
# ============================================================

CHROMA_PATH = "./chroma_db"

chroma_client = chromadb.PersistentClient(
    path=CHROMA_PATH
)

collection = chroma_client.get_or_create_collection(
    name="cloudnexus_documents"
)


# ============================================================
# TEXT SPLITTER
# ============================================================

splitter = RecursiveCharacterTextSplitter(
    chunk_size=1200,
    chunk_overlap=250,
    separators=[
        "\n\n",
        "\n",
        ". ",
        " ",
        "",
    ],
)


# ============================================================
# PDF
# ============================================================

def _load_pdf(path: Path) -> list[Document]:

    loader = PyPDFLoader(str(path))

    documents = loader.load()

    normalized = []

    for document in documents:

        metadata = dict(document.metadata)

        page = metadata.get("page")

        if page is not None:
            metadata["page"] = int(page) + 1

        normalized.append(
            Document(
                page_content=document.page_content,
                metadata=metadata,
            )
        )

    return normalized


# ============================================================
# DOCX
# ============================================================

def _load_docx(path: Path) -> list[Document]:

    loader = Docx2txtLoader(str(path))

    documents = loader.load()

    normalized = []

    for document in documents:

        normalized.append(
            Document(
                page_content=document.page_content,
                metadata=dict(document.metadata),
            )
        )

    return normalized


# ============================================================
# PPTX
# ============================================================

def _load_pptx(path: Path) -> list[Document]:
    """
    Extract PPTX slide-by-slide.

    This is intentionally implemented using
    python-pptx instead of UnstructuredPowerPointLoader
    so that slide numbers remain reliable.
    """

    try:
        from pptx import Presentation
    except ImportError:

        raise RuntimeError(
            "python-pptx is required for PPTX indexing. "
            "Install it with: pip install python-pptx"
        )

    presentation = Presentation(str(path))

    documents = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):

        texts = []

        for shape in slide.shapes:

            if not hasattr(shape, "text"):
                continue

            text = shape.text.strip()

            if text:
                texts.append(text)

        slide_text = "\n".join(texts).strip()

        if not slide_text:
            continue

        documents.append(
            Document(
                page_content=slide_text,
                metadata={
                    "slide": slide_number,
                },
            )
        )

    return documents


# ============================================================
# TXT / MD
# ============================================================

def _load_text(path: Path) -> list[Document]:

    loader = TextLoader(
        str(path),
        encoding="utf-8",
    )

    documents = loader.load()

    return [
        Document(
            page_content=document.page_content,
            metadata=dict(document.metadata),
        )
        for document in documents
    ]


# ============================================================
# LOAD DOCUMENT
# ============================================================

def _load_document(
    local_path: str,
) -> list[Document]:

    path = Path(local_path)

    extension = path.suffix.lower()

    if extension == ".pdf":

        documents = _load_pdf(path)

    elif extension == ".docx":

        documents = _load_docx(path)

    elif extension == ".pptx":

        documents = _load_pptx(path)

    elif extension in {
        ".txt",
        ".md",
        ".markdown",
    }:

        documents = _load_text(path)

    else:

        raise ValueError(
            f"Unsupported file type: {extension}"
        )

    if not documents:

        raise ValueError(
            "No text could be extracted from file."
        )

    return documents


# ============================================================
# NORMALIZE + CHUNK
# ============================================================
def _clean_chunk_text(text: str) -> str:
    """
    Clean extracted document text without
    destroying useful formatting.
    """

    if not text:
        return ""

    # Normalize Windows line endings
    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    # Remove excessive spaces
    text = re.sub(
        r"[ \t]+",
        " ",
        text,
    )

    # Remove excessive blank lines
    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text,
    )

    return text.strip()

def _chunk_pdf_documents(
    documents: list[Document],
) -> list[Document]:
    """
    Chunk PDFs page-by-page so chunks never cross
    unrelated PDF pages.
    """

    chunks = []

    for document in documents:

        page_chunks = splitter.split_documents(
            [document]
        )

        for chunk in page_chunks:

            metadata = dict(
                document.metadata
            )

            chunks.append(
                Document(
                    page_content=_clean_chunk_text(
                        chunk.page_content
                    ),
                    metadata=metadata,
                )
            )

    return chunks

def _chunk_pptx_documents(
    documents: list[Document],
) -> list[Document]:
    """
    Keep PPTX chunks within their original slide.
    """

    chunks = []

    for document in documents:

        slide_chunks = splitter.split_documents(
            [document]
        )

        for chunk in slide_chunks:

            metadata = dict(
                document.metadata
            )

            chunks.append(
                Document(
                    page_content=_clean_chunk_text(
                        chunk.page_content
                    ),
                    metadata=metadata,
                )
            )

    return chunks

def _is_relevant_result(
    score: float,
    query: str,
    text: str,
) -> bool:
    """
    Reject clearly unrelated chunks.

    A chunk is accepted when:
    - it has meaningful keyword overlap, OR
    - it has strong fuzzy similarity.

    RRF score alone is NOT enough because RRF
    always produces a ranking even for unrelated queries.
    """

    query_tokens = set(
        _tokenize(query)
    )

    text_tokens = set(
        _tokenize(text)
    )

    if not query_tokens or not text_tokens:
        return False

    # -----------------------------------------
    # Remove generic question words
    # -----------------------------------------

    stop_words = {
        "what",
        "are",
        "is",
        "the",
        "a",
        "an",
        "of",
        "to",
        "in",
        "on",
        "for",
        "and",
        "or",
        "how",
        "why",
        "when",
        "where",
        "which",
        "who",
        "does",
        "do",
        "this",
        "that",
        "these",
        "those",
    }

    meaningful_query_tokens = (
        query_tokens - stop_words
    )

    if not meaningful_query_tokens:
        return False

    # -----------------------------------------
    # Keyword overlap
    # -----------------------------------------

    overlap = (
        len(
            meaningful_query_tokens
            & text_tokens
        )
        / len(
            meaningful_query_tokens
        )
    )

    # -----------------------------------------
    # Fuzzy similarity
    # -----------------------------------------

    fuzzy_score = partial_ratio(
        query.lower(),
        text.lower(),
    ) / 100.0

    # -----------------------------------------
    # Strong keyword match
    # -----------------------------------------

    if overlap >= 0.30:
        return True

    # -----------------------------------------
    # Strong phrase similarity
    # -----------------------------------------

    if fuzzy_score >= 0.65:
        return True

    return False

def _prepare_chunks(
    local_path: str,
) -> list[Document]:

    documents = _load_document(
        local_path
    )

    cleaned_documents = []

    for document in documents:

        cleaned_text = _clean_chunk_text(
            document.page_content
        )

        if not cleaned_text:
            continue

        cleaned_documents.append(
            Document(
                page_content=cleaned_text,
                metadata=dict(
                    document.metadata
                ),
            )
        )

    if not cleaned_documents:

        raise ValueError(
            "No usable text was extracted."
        )

    extension = Path(
        local_path
    ).suffix.lower()

    # -----------------------------------------
    # PDF
    # -----------------------------------------

    if extension == ".pdf":

        chunks = _chunk_pdf_documents(
            cleaned_documents
        )

    # -----------------------------------------
    # PPTX
    # -----------------------------------------

    elif extension == ".pptx":

        chunks = _chunk_pptx_documents(
            cleaned_documents
        )

    # -----------------------------------------
    # DOCX / TXT / Markdown
    # -----------------------------------------

    else:

        chunks = splitter.split_documents(
            cleaned_documents
        )

        chunks = [
            Document(
                page_content=_clean_chunk_text(
                    chunk.page_content
                ),
                metadata=dict(
                    chunk.metadata
                ),
            )
            for chunk in chunks
            if _clean_chunk_text(
                chunk.page_content
            )
        ]

    if not chunks:

        raise ValueError(
            "No text chunks were created."
        )

    # -----------------------------------------
    # Add sequential chunk indexes
    # -----------------------------------------

    final_chunks = []

    for index, chunk in enumerate(
        chunks
    ):

        metadata = dict(
            chunk.metadata
        )

        metadata[
            "chunk_index"
        ] = index

        final_chunks.append(
            Document(
                page_content=chunk.page_content,
                metadata=metadata,
            )
        )

    return final_chunks

# ============================================================
# INDEX FILE
# ============================================================

def index_file(
    file_id: str,
    local_path: str,
    mime: str,
    folder_id: str | None = None,
    original_filename: str | None = None,
) -> None:

    path = Path(local_path)

    filename = (original_filename if original_filename else path.name)

    print(
        f"\n[RAG] Indexing: {filename}"
    )

    # -----------------------------------------
    # Prepare chunks
    # -----------------------------------------

    chunks = _prepare_chunks(
        local_path
    )

    print(
        f"[RAG] Extracted {len(chunks)} chunks"
    )

    # -----------------------------------------
    # Delete old vectors
    # -----------------------------------------

    delete_file_vectors(
        file_id
    )

    # -----------------------------------------
    # Prepare text
    # -----------------------------------------

    texts = [
        chunk.page_content
        for chunk in chunks
    ]

    # -----------------------------------------
    # Batch embeddings
    # -----------------------------------------

    print(
        "[RAG] Generating embeddings..."
    )

    vectors = embeddings.embed_documents(
        texts
    )

    # -----------------------------------------
    # Prepare Chroma records
    # -----------------------------------------

    ids = []

    metadatas = []

    for index, chunk in enumerate(
        chunks
    ):

        chunk_id = (
            f"{file_id}_{index}"
        )

        metadata = {
            "file_id": file_id,
            "filename": filename,
            "mime": mime,
            "folder_id": folder_id or "",
            "chunk_index": index,
            "source_type": path.suffix.lower().lstrip("."),
        }

        # -------------------------------------
        # Preserve source location
        # -------------------------------------

        original_metadata = (
            chunk.metadata or {}
        )

        if "page" in original_metadata:

            metadata["page"] = str(
                original_metadata["page"]
            )

        if "slide" in original_metadata:

            metadata["slide"] = str(
                original_metadata["slide"]
            )

        # -------------------------------------
        # Preserve other useful metadata
        # -------------------------------------

        for key, value in original_metadata.items():

            if value is None:
                continue

            key = str(key)

            if key in metadata:
                continue

            # Chroma metadata values must be
            # primitive types.
            if isinstance(
                value,
                (str, int, float, bool),
            ):

                metadata[key] = value

            else:

                metadata[key] = str(value)

        ids.append(chunk_id)

        metadatas.append(metadata)

    # -----------------------------------------
    # Store in Chroma
    # -----------------------------------------

    collection.upsert(
        ids=ids,
        documents=texts,
        metadatas=metadatas,
        embeddings=vectors,
    )

    print(
        f"[RAG] Indexed {filename}: "
        f"{len(chunks)} chunks"
    )


# ============================================================
# DELETE FILE VECTORS
# ============================================================

def delete_file_vectors(
    file_id: str,
) -> None:

    results = collection.get(
        where={
            "file_id": file_id
        }
    )

    ids = results.get(
        "ids",
        [],
    )

    if ids:

        collection.delete(
            ids=ids
        )

        print(
            f"[RAG] Deleted {len(ids)} "
            f"vectors for {file_id}"
        )


# ============================================================
# SEMANTIC SEARCH
# ============================================================

def _semantic_search(
    query: str,
    file_ids: list[str] | None,
    folder_id: str | None,
    top_k: int,
) -> list[dict[str, Any]]:

    query_vector = embeddings.embed_query(
        query
    )

    # -----------------------------------------
    # Scope filter
    # -----------------------------------------

    where = None

    if file_ids:

        where = {
            "file_id": {
                "$in": file_ids
            }
        }

    elif folder_id:

        where = {
            "folder_id": folder_id
        }

    # -----------------------------------------
    # Chroma query
    # -----------------------------------------

    kwargs = {
        "query_embeddings": [
            query_vector
        ],
        "n_results": top_k,
    }

    if where:

        kwargs["where"] = where

    results = collection.query(
        **kwargs
    )

    output = []

    result_ids = results.get(
        "ids",
        [],
    )

    if not result_ids:
        return output

    documents = results.get(
        "documents",
        [[]],
    )[0]

    metadatas = results.get(
        "metadatas",
        [[]],
    )[0]

    distances = results.get(
        "distances",
        [[]],
    )[0]

    for index, doc_id in enumerate(
        result_ids[0]
    ):

        metadata = (
            metadatas[index]
            if index < len(metadatas)
            else {}
        )

        output.append(
            {
                "id": doc_id,
                "text": documents[index],
                "file_name": metadata.get(
                    "filename",
                    "Unknown",
                ),
                "metadata": metadata,
                "distance": (
                    distances[index]
                    if index < len(distances)
                    else None
                ),
            }
        )

    return output


# ============================================================
# GET DOCUMENTS FOR BM25
# ============================================================

def _get_documents(
    file_ids: list[str] | None,
    folder_id: str | None,
) -> list[dict[str, Any]]:

    where = None

    if file_ids:

        where = {
            "file_id": {
                "$in": file_ids
            }
        }

    elif folder_id:

        where = {
            "folder_id": folder_id
        }

    kwargs = {
        "include": [
            "documents",
            "metadatas",
        ]
    }

    if where:

        kwargs["where"] = where

    results = collection.get(
        **kwargs
    )

    documents = []

    ids = results.get(
        "ids",
        [],
    )

    texts = results.get(
        "documents",
        [],
    )

    metadatas = results.get(
        "metadatas",
        [],
    )

    for index, doc_id in enumerate(
        ids
    ):

        documents.append(
            {
                "id": doc_id,
                "text": texts[index],
                "metadata": metadatas[index],
            }
        )

    return documents


# ============================================================
# BM25 SEARCH
# ============================================================
def _tokenize(text: str) -> list[str]:
    """
    Better BM25 tokenizer.

    Converts:
        "What are the objectives?"

    into:
        ["what", "are", "the", "objectives"]
    """

    return re.findall(
        r"[a-zA-Z0-9]+",
        text.lower(),
    )


def _expand_query(query: str) -> str:
    """
    Small query expansion layer.

    Helps keyword retrieval handle common
    variations such as:
        objective / objectives
        aim / aims
        method / methods
        result / results
    """

    query = query.strip()

    expansions = {
        "objective": "objectives aim aims goals",
        "objectives": "objective aim aims goals",
        "aim": "aims objective objectives goals",
        "aims": "aim objective objectives goals",
        "goal": "goals objective objectives aim aims",
        "goals": "goal objective objectives aim aims",

        "method": "methods methodology approach",
        "methods": "method methodology approach",
        "result": "results findings outcome",
        "results": "result findings outcome",

        "advantage": "advantages benefit benefits",
        "advantages": "advantage benefit benefits",

        "disadvantage": "disadvantages limitation limitations",
        "disadvantages": "disadvantage limitation limitations",
    }

    tokens = _tokenize(query)

    expanded = [query]

    for token in tokens:

        if token in expansions:
            expanded.append(
                expansions[token]
            )

    return " ".join(expanded)


def _keyword_search(
    query: str,
    documents: list[dict[str, Any]],
    top_k: int,
) -> list[dict[str, Any]]:

    if not documents:
        return []

    tokenized_documents = [
        _tokenize(doc["text"])
        for doc in documents
    ]

    bm25 = BM25Okapi(
        tokenized_documents
    )

    expanded_query = _expand_query(
        query
    )

    query_tokens = _tokenize(
        expanded_query
    )

    scores = bm25.get_scores(
        query_tokens
    )

    ranked_indices = sorted(
        range(len(scores)),
        key=lambda index: scores[index],
        reverse=True,
    )

    results = []

    for index in ranked_indices[:top_k]:

        result = documents[index].copy()

        result["bm25_score"] = float(
            scores[index]
        )

        results.append(result)

    return results


# ============================================================
# HYBRID RETRIEVAL
# ============================================================
def _get_neighbor_chunks(
    result: dict[str, Any],
    all_documents: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """
    Return the retrieved chunk plus its immediate
    neighboring chunks from the same file.

    Example:

        chunk 10 ← previous
        chunk 11 ← retrieved
        chunk 12 ← next
    """

    metadata = result.get(
        "metadata",
        {},
    ) or {}

    file_id = metadata.get(
        "file_id"
    )

    chunk_index = metadata.get(
        "chunk_index"
    )

    if file_id is None or chunk_index is None:
        return [result]

    try:
        chunk_index = int(
            chunk_index
        )
    except (TypeError, ValueError):
        return [result]

    # -----------------------------------------
    # Find neighboring chunks
    # -----------------------------------------

    neighbors = []

    for document in all_documents:

        document_metadata = document.get(
            "metadata",
            {},
        ) or {}

        if document_metadata.get(
            "file_id"
        ) != file_id:
            continue

        document_index = document_metadata.get(
            "chunk_index"
        )

        try:
            document_index = int(
                document_index
            )
        except (TypeError, ValueError):
            continue

        if abs(
            document_index - chunk_index
        ) <= 1:

            neighbors.append(
                document
            )

    # -----------------------------------------
    # Sort by original chunk position
    # -----------------------------------------

    neighbors.sort(
        key=lambda document: int(
            document.get(
                "metadata",
                {},
            ).get(
                "chunk_index",
                0,
            )
        )
    )

    return neighbors
def retrieve(
    query: str,
    folder_id: str | None = None,
    file_ids: list[str] | None = None,
    top_k: int = 8,
) -> list[dict[str, Any]]:

    candidate_k = max(
        top_k * 3,
        20,
    )

    # -----------------------------------------
    # Semantic search
    # -----------------------------------------

    semantic_results = _semantic_search(
        query=query,
        file_ids=file_ids,
        folder_id=folder_id,
        top_k=candidate_k,
    )

    # -----------------------------------------
    # BM25 search
    # -----------------------------------------

    documents = _get_documents(
        file_ids=file_ids,
        folder_id=folder_id,
    )

    keyword_results = _keyword_search(
        query=query,
        documents=documents,
        top_k=candidate_k,
    )

    # -----------------------------------------
    # RRF
    # -----------------------------------------

    rrf_scores = {}
    result_data = {}

    rrf_k = 60

    for rank, result in enumerate(
        semantic_results
    ):

        doc_id = result["id"]

        rrf_scores[doc_id] = (
            rrf_scores.get(
                doc_id,
                0,
            )
            + 1 / (
                rrf_k
                + rank
                + 1
            )
        )

        result_data[doc_id] = result

    for rank, result in enumerate(
        keyword_results
    ):

        doc_id = result["id"]

        rrf_scores[doc_id] = (
            rrf_scores.get(
                doc_id,
                0,
            )
            + 1 / (
                rrf_k
                + rank
                + 1
            )
        )

        if doc_id not in result_data:

            metadata = result.get(
                "metadata",
                {},
            )

            result_data[doc_id] = {
                "id": doc_id,
                "text": result["text"],
                "file_name": metadata.get(
                    "filename",
                    "Unknown",
                ),
                "metadata": metadata,
            }

    # -----------------------------------------
    # RRF ranking
    # -----------------------------------------

    ranked = sorted(
        rrf_scores.items(),
        key=lambda item: item[1],
        reverse=True,
    )

    # -----------------------------------------
    # Relevance reranking
    # -----------------------------------------

    query_tokens = set(
        _tokenize(query)
    )

    reranked = []

    for doc_id, rrf_score in ranked:

        result = result_data[doc_id]

        text = result.get(
            "text",
            "",
        )

        text_tokens = set(
            _tokenize(text)
        )

        overlap = 0.0

        if query_tokens and text_tokens:

            overlap = (
                len(
                    query_tokens
                    & text_tokens
                )
                / len(query_tokens)
            )

        fuzzy_score = partial_ratio(
            query.lower(),
            text.lower(),
        ) / 100.0

        final_score = (
            (rrf_score * 0.65)
            + (overlap * 0.20)
            + (fuzzy_score * 0.15)
        )

        reranked.append(
            (
                doc_id,
                final_score,
            )
        )

    ranked = sorted(
        reranked,
        key=lambda item: item[1],
        reverse=True,
    )

    # -----------------------------------------
    # Build filtered results
    # -----------------------------------------

    final_results = []

    for doc_id, score in ranked:

        if len(final_results) >= top_k:
            break

        result = result_data[
            doc_id
        ].copy()

        text = result.get(
            "text",
            "",
        )

        # -------------------------------------
        # Confidence filter
        # -------------------------------------

        if not _is_relevant_result(
            score,
            query,
            text,
        ):
            continue

        result["hybrid_score"] = float(
            score
        )

        metadata = result.get(
            "metadata",
            {},
        ) or {}

        result["file_id"] = metadata.get(
            "file_id",
            result.get("file_id"),
        )

        result["file_name"] = metadata.get(
            "filename",
            result.get(
                "file_name",
                "Unknown",
            ),
        )

        result["source_type"] = metadata.get(
            "source_type",
            result.get("source_type"),
        )

        if "page" in metadata:
            result["page"] = metadata[
                "page"
            ]

        if "slide" in metadata:
            result["slide"] = metadata[
                "slide"
            ]

        result["metadata"] = metadata

        # -------------------------------------
        # Preserve original citation text
        # -------------------------------------

        result["retrieved_text"] = text

        final_results.append(
            result
        )

    # -----------------------------------------
    # Context expansion
    # -----------------------------------------

    expanded_results = []

    for result in final_results:

        neighbors = _get_neighbor_chunks(
            result,
            documents,
        )

        result_copy = result.copy()

        result_copy["retrieved_text"] = (
            result.get(
                "retrieved_text",
                result.get(
                    "text",
                    "",
                ),
            )
        )

        context_parts = []

        for neighbor in neighbors:

            neighbor_text = neighbor.get(
                "text",
                "",
            ).strip()

            if neighbor_text:

                context_parts.append(
                    neighbor_text
                )

        result_copy["context_text"] = (
            "\n\n".join(
                context_parts
            )
        )

        result_copy["context_chunks"] = len(
            neighbors
        )

        expanded_results.append(
            result_copy
        )

    return expanded_results